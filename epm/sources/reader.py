"""Read CSV and XLSX files into a grid of parsed cell values."""

from __future__ import annotations

import csv
import datetime as dt
import io
import math
import re
from dataclasses import dataclass, field
from typing import Any

from .a1 import cell_name

_MISSING = {"", "n/a", "na", "#n/a", "nan", "null", "none", "#value!", "#div/0!", "#ref!"}
_DASHES = {"-", "–", "—", "--"}
_NUMBER = re.compile(r"^[+-]?(\d+(\.\d*)?|\.\d+)([eE][+-]?\d+)?$")


class SourceError(Exception):
    def __init__(self, code: str, message: str, **detail: Any):
        super().__init__(message)
        self.code = code
        self.message = message
        self.detail = detail


def parse_value(raw: Any) -> Any:
    """Turn a raw cell into a float, ``None`` (blank), or the original value if not numeric.

    Text cells accept thousands separators, currency symbols, ``(123)`` for negatives and a
    trailing ``%``. A lone dash is zero, as in accounting formats.
    """
    if raw is None:
        return None
    if isinstance(raw, bool):
        return float(raw)
    if isinstance(raw, int | float):
        f = float(raw)
        return None if math.isnan(f) else f
    if not isinstance(raw, str):
        return raw
    text = raw.strip()
    if text.lower() in _MISSING:
        return None
    if text in _DASHES:
        return 0.0
    body = text.replace(",", "").replace(" ", "").replace(" ", "")
    negative = False
    if body.startswith("(") and body.endswith(")"):
        negative, body = True, body[1:-1]
    for symbol in ("$", "€", "£", "¥"):
        body = body.replace(symbol, "")
    percent = body.endswith("%")
    if percent:
        body = body[:-1]
    if not _NUMBER.match(body):
        return raw.strip()
    value = float(body)
    if percent:
        value /= 100
    return -value if negative else value


def label_of(raw: Any) -> str | None:
    """Normalize a label cell: ``2024.0`` → ``"2024"``, dates → ISO dates, text stripped."""
    if raw is None:
        return None
    if isinstance(raw, dt.datetime):
        return raw.date().isoformat() if raw.time() == dt.time() else raw.isoformat()
    if isinstance(raw, dt.date):
        return raw.isoformat()
    if isinstance(raw, float) and raw.is_integer():
        return str(int(raw))
    text = str(raw).strip()
    if not text:
        return None
    try:
        f = float(text)
    except ValueError:
        return text
    return str(int(f)) if f.is_integer() and "e" not in text.lower() else text


@dataclass
class Sheet:
    """Cells of one sheet (or CSV file), addressed with 1-based row and column numbers."""

    name: str
    rows: list[list[Any]]
    raw: list[list[Any]]
    uncached_formulas: set[tuple[int, int]] = field(default_factory=set)

    @property
    def n_rows(self) -> int:
        return len(self.rows)

    @property
    def n_cols(self) -> int:
        return max((len(r) for r in self.rows), default=0)

    def get(self, row: int, col: int) -> Any:
        """Parsed value at (row, col), or ``None`` outside the sheet."""
        if 1 <= row <= len(self.rows) and 1 <= col <= len(self.rows[row - 1]):
            return self.rows[row - 1][col - 1]
        return None

    def get_raw(self, row: int, col: int) -> Any:
        if 1 <= row <= len(self.raw) and 1 <= col <= len(self.raw[row - 1]):
            return self.raw[row - 1][col - 1]
        return None

    def ref(self, row: int, col: int) -> str:
        return f"{self.name}!{cell_name(row, col)}"


def read_csv(data: bytes, name: str, delimiter: str = ",") -> Sheet:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    raw = [list(r) for r in csv.reader(io.StringIO(text), delimiter=delimiter)]
    return Sheet(name, [[parse_value(c) for c in r] for r in raw], raw)


def read_xlsx(data: bytes, file: str) -> dict[str, Sheet]:
    """Read every sheet. Formula cells use the values Excel last saved."""
    import openpyxl

    try:
        values = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        formulas = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    except Exception as e:  # noqa: BLE001 - openpyxl raises many types for bad files
        raise SourceError("unreadable_file", f"cannot read '{file}' as xlsx: {e}") from None
    sheets = {}
    for ws in values.worksheets:
        raw = [list(r) for r in ws.iter_rows(values_only=True)]
        uncached = set()
        for r_i, row in enumerate(formulas[ws.title].iter_rows(values_only=True), start=1):
            for c_i, v in enumerate(row, start=1):
                is_formula = isinstance(v, str) and v.startswith("=")
                if is_formula and (r_i > len(raw) or raw[r_i - 1][c_i - 1] is None):
                    uncached.add((r_i, c_i))
        sheets[ws.title] = Sheet(
            ws.title, [[parse_value(c) for c in r] for r in raw], raw, uncached
        )
    return sheets
