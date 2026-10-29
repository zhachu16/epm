"""Write inputs/historicals.xlsx for the sample LBO (Project Atlas, a fictional company).

Run from anywhere: python examples/lbo/make_inputs.py
"""

from pathlib import Path

import openpyxl
from openpyxl.styles import Font

HERE = Path(__file__).parent
YEARS = [2022, 2023, 2024]

INCOME_STATEMENT = [
    ("Revenue", [820.0, 900.0, 1000.0]),
    ("Cost of goods sold", [492.0, 531.0, 580.0]),
    ("SG&A", [164.0, 175.0, 190.0]),
    ("Depreciation & amortization", [33.0, 36.0, 40.0]),
    ("Interest expense", [20.0, 19.0, 18.0]),
    ("Income tax", [27.75, 34.75, 43.0]),
]

BALANCE_SHEET = [
    ("Cash", [60.0, 70.0, 80.0]),
    ("Accounts receivable", [112.0, 123.0, 137.0]),
    ("Inventory", [81.0, 87.0, 95.0]),
    ("Accounts payable", [54.0, 58.0, 64.0]),
    ("PP&E, net", [300.0, 320.0, 345.0]),
    ("Existing debt", [250.0, 240.0, 230.0]),
]

DEAL = [
    ("Entry multiple (x LTM EBITDA)", 9.0),
    ("Transaction fees (% of EV)", 0.02),
    ("Financing fees (% of debt raised)", 0.025),
]


def sheet(wb, title, rows):
    ws = wb.create_sheet(title)
    ws.append(["USD m", *YEARS])
    for label, values in rows:
        ws.append([label, *values])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 30


def main() -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheet(wb, "IS", INCOME_STATEMENT)
    sheet(wb, "BS", BALANCE_SHEET)
    deal = wb.create_sheet("Deal")
    deal.append(["Assumption", "Value"])
    for label, value in DEAL:
        deal.append([label, value])
    deal.column_dimensions["A"].width = 36
    out = HERE / "inputs" / "historicals.xlsx"
    out.parent.mkdir(exist_ok=True)
    wb.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
